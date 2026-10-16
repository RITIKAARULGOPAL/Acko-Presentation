"""List and export picture shapes from the proposal deck (slide -> images) as WebP."""
import sys, os, io, json
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation
from PIL import Image

src, out = sys.argv[1], sys.argv[2]
os.makedirs(out, exist_ok=True)
prs = Presentation(src)
W, H = prs.slide_width, prs.slide_height
index = []

def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from walk(sh.shapes)
        elif sh.shape_type == 13:
            yield sh

for si, slide in enumerate(prs.slides, 1):
    for k, sh in enumerate(walk(slide.shapes)):
        try:
            blob = sh.image.blob
        except Exception:
            continue
        im = Image.open(io.BytesIO(blob))
        cover = (sh.width * sh.height) / (W * H)
        name = f"s{si:02d}_{k:02d}"
        im = im.convert("RGBA") if im.mode in ("P", "LA", "RGBA") else im.convert("RGB")
        if max(im.size) > 2560:
            im.thumbnail((2560, 2560))
        im.save(os.path.join(out, name + ".webp"), "WEBP", quality=82, method=6)
        index.append(dict(name=name, slide=si, w=im.size[0], h=im.size[1], cover=round(cover, 3),
                          x=round(sh.left / W, 3), y=round(sh.top / H, 3)))
json.dump(index, open(os.path.join(out, "_index.json"), "w"), indent=1)
print(len(index), "images")
for r in index:
    if r["cover"] > 0.02:
        print(r)
