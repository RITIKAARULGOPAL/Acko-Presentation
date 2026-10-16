"""Pull the About SKV images (deck slides 3-13) out of the proposal PPTX into assets/img/skv/.

Usage (from the project folder):
  python tools/extract_skv.py "<proposal.pptx>" [out_dir]

The deck's pictures are raster fills on freeform shapes, often inside groups and cropped
with srcRect / fillRect, so each image is cut exactly as it appears on the slide.
Shapes are picked by slide number + Google-Slides shape id (stable for this deck).
"""
import io, os, sys
from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)

# (slide, shape id) -> (output name, max px on the long edge)
PICK = {
    (5, "337"): ("apart", 1100),
    **{(5, sid): (f"icon-apart-{k}", 160) for k, sid in enumerate(["356", "357", "358", "359", "360", "361"], 1)},
    (6, "377"): ("icon-svc-civil", 160), (6, "379"): ("icon-svc-mep", 160),
    (6, "381"): ("icon-svc-furn", 160), (6, "383"): ("icon-svc-it", 160),
    (6, "384"): ("svc-1", 1000), (6, "386"): ("svc-2", 1400), (6, "385"): ("svc-3", 800),
    (7, "401"): ("stand-1", 1000), (7, "402"): ("stand-2", 1400), (7, "403"): ("stand-3", 800),
    (9, "527"): ("factory", 2000),
    (9, "546"): ("icon-fac-area", 160), (9, "531"): ("icon-fac-team", 160),
    (9, "538"): ("icon-fac-crew", 160), (9, "539"): ("icon-fac-plant", 160),
    (10, "565"): ("cutaway", 2000),
    (10, "578"): ("icon-wood", 160), (10, "579"): ("icon-metal", 160),
    (10, "580"): ("icon-paint", 160), (10, "581"): ("icon-uph", 160),
    (11, "601"): ("create-1", 900), (11, "602"): ("create-2", 1200),
    (11, "599"): ("create-3", 1100), (11, "600"): ("create-4", 1100),
    (12, "635"): ("team-tushar", 480), (12, "655"): ("team-swati", 480), (12, "632"): ("team-rishabh", 480),
    (12, "622"): ("team-ashish", 480), (12, "626"): ("team-atul", 480), (12, "665"): ("team-jegan", 480),
    (12, "667"): ("team-umar", 480), (12, "669"): ("team-bisa", 480), (12, "670"): ("team-nabeela", 480),
    (12, "672"): ("team-ishwarya", 480), (12, "671"): ("team-sai", 480), (12, "674"): ("team-parinay", 480),
    **{(13, sid): (f"project-{k:02d}", 700) for k, sid in enumerate(["680", "679", "681", "682", "683", "684", "703", "702", "701", "688"], 1)},
    (13, "696"): ("t-jci", 260), (13, "697"): ("t-levis", 200), (13, "698"): ("t-gril", 260),
}
# slide 7: one tall PNG per column holding a stack of icons -> sliced into single icons
STRIPS = {(7, "398"): "icon-quality", (7, "399"): "icon-sustain", (7, "400"): "icon-safety"}
# images that sit on a white background in the deck: white becomes transparent
MATTE = {"team-atul"}
# cut-outs with a faint dark halo left over from background removal: drop the semi-transparent fringe
DEHALO = {"team-ishwarya"}
LOGO_SLIDE = 8
# (slide 3's animated ASCII pattern is redrawn live on a canvas in the page; the GIF would be ~3-10 MB)


def crop_frac(im, x0, y0, x1, y1):
    """Crop by fractions of the image; ranges outside 0..1 pad with transparency."""
    W, H = im.size
    box = [round(x0 * W), round(y0 * H), round(x1 * W), round(y1 * H)]
    if box == [0, 0, W, H]:
        return im
    out = Image.new("RGBA", (max(1, box[2] - box[0]), max(1, box[3] - box[1])), (0, 0, 0, 0))
    out.paste(im, (-box[0], -box[1]))
    return out


def visible(im, blipfill):
    """Apply srcRect (crop/pad) then stretch/fillRect (image scaled past the shape) -> what the slide shows."""
    f = lambda el, k: int(el.get(k, 0)) / 100000 if el is not None else 0
    src = blipfill.find(qn("a:srcRect"))
    im = crop_frac(im, f(src, "l"), f(src, "t"), 1 - f(src, "r"), 1 - f(src, "b"))
    st = blipfill.find(qn("a:stretch"))
    fr = st.find(qn("a:fillRect")) if st is not None else None
    if fr is not None:
        l, t, r, b = f(fr, "l"), f(fr, "t"), f(fr, "r"), f(fr, "b")
        sw, sh = 1 - r - l, 1 - b - t
        im = crop_frac(im, -l / sw, -t / sh, (1 - l) / sw, (1 - t) / sh)
    return im


def shapes(shs, T=(0, 0, 1, 1, 0, 0)):
    """Yield (shape, x, y, w, h) in slide EMU, resolving group transforms."""
    tr = lambda x, y: (T[0] + (x - T[4]) * T[2], T[1] + (y - T[5]) * T[3])
    for sh in shs:
        if sh.shape_type == 6:
            g = sh._element.find(qn("p:grpSpPr")).find(qn("a:xfrm"))
            off, ext, co, ce = (g.find(qn(n)) for n in ("a:off", "a:ext", "a:chOff", "a:chExt"))
            gx, gy = tr(int(off.get("x")), int(off.get("y")))
            yield from shapes(sh.shapes, (gx, gy, T[2] * int(ext.get("cx")) / max(1, int(ce.get("cx"))),
                                          T[3] * int(ext.get("cy")) / max(1, int(ce.get("cy"))), int(co.get("x")), int(co.get("y"))))
            continue
        x, y = tr(sh.left or 0, sh.top or 0)
        yield sh, x, y, (sh.width or 0) * T[2], (sh.height or 0) * T[3]


def image_of(slide, sh):
    blip = sh._element.find(".//" + qn("a:blip"))
    if blip is None:
        return None, None
    part = slide.part.related_part(blip.get(qn("r:embed")))
    return Image.open(io.BytesIO(part.blob)), blip.getparent()


def save(im, path, maxpx):
    im = im.convert("RGBA")
    im.thumbnail((maxpx, maxpx), Image.LANCZOS)
    if im.getextrema()[3][0] == 255:  # fully opaque
        im = im.convert("RGB")
    im.save(path, "WEBP", quality=84, method=6)


def unwhite(im, hard=False):
    """White -> transparent. hard: only near-pure white (photos); else alpha follows ink density (line icons)."""
    im = im.convert("RGBA")
    px = im.load()
    for y in range(im.height):
        for x in range(im.width):
            r, g, b, a = px[x, y]
            if hard:
                if min(r, g, b) >= 246:
                    px[x, y] = (r, g, b, 0)
            else:
                ink = 255 - min(r, g, b)
                px[x, y] = (r, g, b, min(a, min(255, ink * 2)))
    return im


def slice_strip(im, name, out):
    """Split a vertical stack of icons (drawn on white) into separate transparent icons."""
    im = unwhite(im)
    a = im.split()[3]
    rows = [any(a.getpixel((x, y)) > 40 for x in range(a.width)) for y in range(a.height)]
    bands, y0 = [], None
    for y, on in enumerate(rows + [False]):
        if on and y0 is None:
            y0 = y
        elif not on and y0 is not None:
            if y - y0 > a.height * 0.04:
                bands.append((y0, y))
            y0 = None
    for k, (t, b) in enumerate(bands, 1):
        save(im.crop((0, t, im.width, b)), os.path.join(out, f"{name}-{k}.webp"), 160)
    return len(bands)


def main():
    deck = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(ROOT, "assets", "img", "skv")
    os.makedirs(out, exist_ok=True)
    prs = Presentation(deck)
    done, logos = set(), []
    for n, slide in enumerate(prs.slides, 1):
        if not 3 <= n <= 13:
            continue
        for sh, x, y, w, h in shapes(slide.shapes):
            sid = sh.name.split(";")[1] if ";" in sh.name else sh.name
            src, bf = image_of(slide, sh)
            if src is None:
                continue
            if n == LOGO_SLIDE:
                logos.append((y + h / 2, x + w / 2, visible(src.convert("RGBA"), bf)))
            elif (n, sid) in STRIPS:
                k = slice_strip(visible(src.convert("RGBA"), bf), STRIPS[(n, sid)], out)
                print(f"  s{n:02d} {STRIPS[(n, sid)]} x{k}")
                done.add((n, sid))
            elif (n, sid) in PICK:
                name, maxpx = PICK[(n, sid)]
                im = visible(src.convert("RGBA"), bf)
                if name in MATTE:
                    im = unwhite(im, hard=True)
                if name in DEHALO:
                    im.putalpha(im.split()[3].point(lambda a: 255 if a >= 200 else 0))
                save(im, os.path.join(out, name + ".webp"), maxpx)
                print(f"  s{n:02d} {name}")
                done.add((n, sid))
    # logo wall: cluster into rows by vertical centre, then left -> right
    logos.sort(key=lambda t: t[0])
    rows, cur = [], [logos[0]]
    for l in logos[1:]:
        if l[0] - cur[-1][0] > prs.slide_height * 0.05:
            rows.append(cur); cur = [l]
        else:
            cur.append(l)
    rows.append(cur)
    k = 0
    for r in rows:
        for _, _, im in sorted(r, key=lambda t: t[1]):
            k += 1
            save(im, os.path.join(out, f"logo-{k:02d}.webp"), 360)
    print(f"  s08 logos: {k} in {len(rows)} rows ({', '.join(str(len(r)) for r in rows)})")
    missing = [v[0] for key, v in PICK.items() if key not in done] + [v for key, v in STRIPS.items() if key not in done]
    if missing:
        print("  MISSING (deck changed?):", ", ".join(missing))
    print("done ->", out)


if __name__ == "__main__":
    main()
