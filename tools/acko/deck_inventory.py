"""List every picture on every slide of the Acko design deck: media file, pixel size, position and size on the slide
(in % of the slide), crop, and the shape's name. Writes tools/_media/inventory.json and prints a compact table.

    python tools/acko/deck_inventory.py "<deck.pptx>" [slide numbers…]
"""
import io, json, os, sys, zipfile
from pptx import Presentation
from pptx.util import Emu
from PIL import Image

deck = sys.argv[1]
only = {int(a) for a in sys.argv[2:]}
p = Presentation(deck)
W, H = p.slide_width, p.slide_height
z = zipfile.ZipFile(deck)
out = {}

def walk(shapes, off=(0, 0, 1, 1)):
    for sh in shapes:
        if sh.shape_type == 6:                       # group: children are in the group's own frame
            gx, gy, sx, sy = off
            ch = sh._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}chOff')
            ce = sh._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}chExt')
            if ch is not None and ce is not None and int(ce.get('cx')) and int(ce.get('cy')):
                kx, ky = sh.width / int(ce.get('cx')), sh.height / int(ce.get('cy'))
                yield from walk(sh.shapes, (gx + sx * (sh.left - int(ch.get('x')) * kx), gy + sy * (sh.top - int(ch.get('y')) * ky), sx * kx, sy * ky))
            else:
                yield from walk(sh.shapes, off)
            continue
        yield sh, off

for i, s in enumerate(p.slides, 1):
    if only and i not in only:
        continue
    rows = []
    for sh, (gx, gy, sx, sy) in walk(s.shapes):
        if sh.shape_type != 13 and not hasattr(sh, 'image'):
            continue
        try:
            img = sh.image
        except Exception:
            continue
        name = os.path.basename(img.filename) if img.filename else '?'
        part = sh.part.related_part(sh._element.blip_rId) if hasattr(sh._element, 'blip_rId') else None
        media = os.path.basename(part.partname) if part is not None else name
        try:
            w, h = Image.open(io.BytesIO(img.blob)).size
        except Exception:
            w = h = 0
        x, y = gx + sx * (sh.left or 0), gy + sy * (sh.top or 0)
        cw, chh = sx * (sh.width or 0), sy * (sh.height or 0)
        crop = [round(getattr(sh, k, 0) * 100, 1) for k in ('crop_left', 'crop_top', 'crop_right', 'crop_bottom')]
        rows.append(dict(media=media, px=[w, h], kb=len(img.blob) // 1024, name=sh.name,
                         box=[round(x / W * 100, 1), round(y / H * 100, 1), round(cw / W * 100, 1), round(chh / H * 100, 1)], crop=crop))
    out[i] = rows
    print(f"--- slide {i}")
    for r in rows:
        print(f"  {r['media']:<14} {r['px'][0]}x{r['px'][1]:<5} {r['kb']:>6}KB box%={r['box']} crop={r['crop']} {r['name'][:30]}")

os.makedirs('tools/_media', exist_ok=True)
with open('tools/_media/inventory.json', 'w') as f:
    json.dump(out, f, indent=1)
