"""Text boxes per slide with their position (% of the slide), font size and colour: for laying out the HTML like the deck.
    python tools/acko/deck_text.py "<deck.pptx>" 2 5 8 …"""
import sys
from pptx import Presentation
p = Presentation(sys.argv[1]); W, H = p.slide_width, p.slide_height
want = {int(a) for a in sys.argv[2:]}
def walk(shapes, off=(0, 0, 1, 1)):
    for sh in shapes:
        if sh.shape_type == 6:
            gx, gy, sx, sy = off
            A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
            ch, ce = sh._element.find('.//' + A + 'chOff'), sh._element.find('.//' + A + 'chExt')
            if ch is not None and int(ce.get('cx')) and int(ce.get('cy')):
                kx, ky = sh.width / int(ce.get('cx')), sh.height / int(ce.get('cy'))
                yield from walk(sh.shapes, (gx + sx * (sh.left - int(ch.get('x')) * kx), gy + sy * (sh.top - int(ch.get('y')) * ky), sx * kx, sy * ky))
            else: yield from walk(sh.shapes, off)
            continue
        yield sh, off
for i, s in enumerate(p.slides, 1):
    if i not in want: continue
    print(f"--- slide {i}")
    for sh, (gx, gy, sx, sy) in walk(s.shapes):
        if not sh.has_text_frame or not sh.text_frame.text.strip(): continue
        x, y = gx + sx * (sh.left or 0), gy + sy * (sh.top or 0)
        sz = col = font = None
        for par in sh.text_frame.paragraphs:
            for r in par.runs:
                if r.font.size and not sz: sz = r.font.size.pt
                if r.font.name and not font: font = r.font.name
                try:
                    if r.font.color and r.font.color.type is not None and not col: col = str(r.font.color.rgb)
                except Exception: pass
        t = " / ".join(par.text.strip() for par in sh.text_frame.paragraphs if par.text.strip())
        print(f"  [{x/W*100:5.1f},{y/H*100:5.1f} {sx*(sh.width or 0)/W*100:5.1f}x{sy*(sh.height or 0)/H*100:4.1f}] {sz}pt {font} #{col}  {t[:110]}")
